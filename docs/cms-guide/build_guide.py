"""Build the FEKRA CMS manual with the bundled Codex Python runtime."""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

ROOT = Path(__file__).resolve().parent
SECTIONS = [
('Find your way around the CMS','02-dashboard',
 'Use the menu button at the top left to open navigation. Collections hold individual records; the Site area holds settings used across the website.',[
 'Open Pages for the homepage, contact content, and other standalone pages. Open Blog for articles, and Services for service landing pages.',
 'Open Careers for job adverts and Media for reusable images, videos, and PDFs. Authors and Categories organize blog content.',
 'Open Applications or Contact messages to review submissions. Admins also see CV files in the menu.',
 'Open Header & Navigation, Footer, or Site Settings for sitewide changes. Use the top right Locale control to choose the editing language.'
 ],'Editor accounts can manage content and submissions. Admin accounts also manage users and Site Settings.'),
('Create or update a page','04-page-editor',
 'A page contains a title, one shared URL slug, and a sequence of layout blocks. The homepage uses the special slug home.',[
 'Choose Pages, then open an existing title. To create a page, select Create New from the list.',
 'Enter the Title and a short Slug such as about. Keep an existing live slug unless a developer has arranged the redirect.',
 'Open Content and edit the Layout blocks. Use Settings to manage Available Locales and Hide From Sitemap.',
 'Save Draft, review the page with Preview, then choose Publish changes. Open the public URL to confirm the result.'
 ],'Pages require explicit saving. A page is not added to the header or footer automatically when published.'),
('Build a page from sections','21-block-picker',
 'Layout blocks are the sections of a page, shown from top to bottom in the same order as the website.',[
 'Select Add Layout and choose the appropriate block. Hero introduces a page; Rich text holds long copy; CTA band adds a call to action.',
 'Expand the new block using its arrow. Fill required fields marked with an asterisk, choose media, and configure any buttons.',
 'Use the drag handle to move a block. Open its three dot menu for available actions; check the selected block before removing it.',
 'Collapse sections to navigate a long page. Save Draft and review the full composition before publishing.'
 ],'Other blocks include FAQ, Testimonials, Process steps, Card grid, Tech stack, Industry grid, Contact section, and Booking embed.'),
('Edit headings images and buttons','18-block-editor',
 'Each block exposes fields for its content and supported visual options. Start with a small change and check it in context.',[
 'Expand the relevant block. In a Hero block, edit Heading, Heading Accent, and Body. Rotating Words replace the fixed accent line when present.',
 'Select Create New or Choose from existing in a media field. Use descriptive alt text on the media record.',
 'Expand Ctas or the link group. Set a visible Label and choose Internal page, Built-in route, or External URL.',
 'Use an internal document relationship for a CMS page, a built-in route for a listing such as Blog, and a full https:// URL for an external destination.'
 ],'A link without a label is not rendered. Check button destinations and images in Preview at mobile and desktop sizes.'),
('Save preview publish and restore','04-page-editor',
 'Drafts let you review content before changing the public website. The top toolbar shows the document status and publishing controls.',[
 'Save Draft after editing. Blog and Services also autosave drafts, but wait for saving to finish before leaving or translating.',
 'Use the eye icon for Live Preview or the adjacent arrow icon to open a preview. Choose Mobile, Tablet, Laptop, or Desktop when the preview offers device controls.',
 'Choose Publish changes when the copy, layout, links, SEO, and translations are approved. Then check the public page outside preview mode.',
 'To recover older content, open Versions, select the relevant version, compare it, and use the restore control offered there. Review the restored content before publishing.'
 ],'Header, Footer, Site Settings, Media, Authors, and Categories save directly. Their Save action does not use the content draft review process.'),
('Set search and social information','20-seo',
 'Pages, Blog, Services, and Careers have an SEO tab. It controls the search title, description, and social sharing image.',[
 'Open SEO. Enter a clear Title and Description; the CMS displays length guidance and a search result preview.',
 'Choose a Meta Image. A 1200 by 630 image is a useful social sharing size; Site Settings provides the default image when a document has none.',
 'Leave Canonical Override empty for the normal page URL. Ask the site owner before using a different canonical.',
 'Enable Noindex only when the page should be excluded from search results and the sitemap. Save and publish the content change.'
 ],'Hide From Sitemap in page Settings only excludes that page from the sitemap. Noindex also requests exclusion from search results.'),
('Manage languages and shared sections','19-locales',
 'The CMS supports English EN, Arabic AR, German DE, French FR, and Spanish ES. Localized fields carry a language indicator.',[
 'Select the editing language using Locale at the top right. Translate localized copy, including titles, block text, button labels, and SEO fields.',
 'Keep the same slug across languages. English pages use /about; Arabic pages use /ar/about, with equivalent prefixes for the other languages.',
 'Add a language to Available Locales only when its translation is approved. Find this field under Settings for Pages, Meta for Blog, Structure for Services, or Details for Careers.',
 'For a Shared section block, edit its source section on the Home page in the same locale, then publish Home and check the pages that reuse it.'
 ],'Shared fields such as slugs, images, relationships, and layout structure can affect multiple languages. A Shared section reuses homepage content rather than maintaining a separate copy.'),
('Write and publish a blog article','06-blog-create',
 'Blog articles combine an article body with an excerpt, hero image, category, author, optional extra sections, and SEO.',[
 'Open Blog and select Create New. Enter a Title and Slug, a concise Excerpt of up to 320 characters, and a Hero Image.',
 'Keep Content Format set to Rich text editor. Write the body and use H2, H3, or H4 headings. The page template provides the main H1.',
 'Open Meta to choose Category, Author, Tags, Related Posts, and Available Locales. Create author or category records first if needed.',
 'Complete SEO, set Featured if appropriate, review Published At, and preview the article. Publish changes and check its listing on /blog.'
 ],'Authors store name, role, photo, bio, and LinkedIn URL. Categories store a title, slug, and description.'),
('Import an HTML blog article','07-blog-html',
 'Use the custom HTML editor when an article has already been prepared as HTML. The normal article title, image, metadata, and SEO still apply.',[
 'In the article Content tab, choose Import / paste HTML under Content Format.',
 'Paste the markup into HTML content or choose an .html file using the import control. Keep the file smaller than 500 KB.',
 'Use inline CSS or a style block and absolute image URLs. JavaScript and forms are disabled in the rendered article.',
 'Review both Desktop and Mobile under HTML preview. Save the draft, check the complete article preview, then publish.'
 ],'Put article markup in HTML content, not Excerpt. Excerpt is a short plain language summary for article cards.'),
('Translate blog articles','23-translate',
 'The custom Translate empty languages button creates draft translations from the English article for languages that are still empty.',[
 'Select EN and complete the article. Save the English changes first; the translation button is disabled while the form has unsaved changes.',
 'Select Translate empty languages. Allow the operation to finish and read any success or failure message.',
 'Switch to each translated locale and review the title, body, excerpt, and SEO. Existing edited translations are kept by the normal button.',
 'Check Available Locales in Meta. The translation action adds successful languages to the draft list, so retain only languages approved for release before publishing.'
 ],'Translations are saved as drafts and require editorial review. If a language fails, keep the English source saved and retry after checking the reported error.'),
('Create and organize services','22-service-structure',
 'Services have their own landing pages, page builder layout, and navigation structure.',[
 'Open Services and choose Create New. Enter Title, Slug, and required Summary. Choose an Icon and build the page in Content.',
 'Open Structure. Choose Parent for a related child landing page, and select Related Services to connect relevant pages.',
 'Set Order to control the overview sequence; lower numbers appear first. Add Menu Roles when the service should provide role labels in the Services mega menu.',
 'Set Available Locales, complete SEO, and preview. Publish, then check the landing page, overview, breadcrumbs, and header menu.'
 ],'The Parent field supports service relationships and breadcrumbs. Keep live slugs stable even when reorganizing services.'),
('Publish and close job vacancies','10-job-details',
 'A job has two separate states: its CMS draft or published state, and its open or closed recruitment status.',[
 'Open Careers and choose Create New. In Role, enter Title, Slug, required Summary and Description, plus Requirements and Benefits.',
 'In Details, enter Location and choose Work Model and Employment Type. Review Department, Country Code, City, Valid Through, and Available Locales.',
 'Keep Role Status set to Open when applications should be accepted. Complete SEO, preview, and Publish changes.',
 'To stop recruitment, set Role Status to Closed and publish that change. The page stays reachable, its application form is disabled, and JobPosting structured data is removed.'
 ],'The Careers collection manages individual job adverts. Page builder content for careers landing sections belongs in Pages where configured.'),
('Upload and reuse media','11-media',
 'Media is the reusable library for public images, MP4 videos, and PDF documents. Applicant CVs belong in the separate private CV files collection.',[
 'Open Media and choose Create New. Select a file, drag it into the upload area, or use the offered URL option.',
 'Write useful Alt text for meaningful images. Choose Decorative only when an image conveys no information.',
 'Add a localized Caption and a Credit where needed. Save the media record, then select it from the relevant page or article field.',
 'For an image, review the focal point and the resulting crop where the upload editor offers those controls. Check both a card and the full page.'
 ],'Media is saved directly and may be reused in several places. Check dependencies before replacing or deleting a shared asset.'),
('Update header links and announcements','12-header',
 'Header & Navigation controls the main menu, nested dropdown links, call to action buttons, and the announcement.',[
 'Open Header & Navigation. Expand the menu item to change, or add a Menu item. Use the drag handle to adjust order.',
 'Set a visible Label and choose the correct link type and destination. Use Children for dropdown entries and add descriptions when useful.',
 'Review Ctas for header buttons. In Announcement, enable the banner, enter localized Text, and configure its link.',
 'Select Save. Test the header links, dropdowns, announcement, and mobile menu on the public site in each affected language.'
 ],'The header supports up to 8 menu items, 12 dropdown items per menu entry, and 2 CTA buttons. Saving makes the change directly.'),
('Maintain the footer','13-footer',
 'The Footer global manages its introductory copy, link columns, newsletter text, legal links, and copyright.',[
 'Open Footer and choose the editing Locale. Update Tagline and Blurb.',
 'Add or edit Columns, their titles, and links. Configure each link label and destination using the same link controls as the header.',
 'Review Newsletter Enabled, Heading, and Body. Add the approved Privacy, Cookies, and Terms destinations under Legal Links.',
 'Update Copyright and Save. Check the public footer on desktop and mobile, including newsletter text and every legal link.'
 ],'The footer supports up to 4 link columns and 6 legal links. Changes to shared link structure can affect multiple languages.'),
('Maintain site identity and contact settings','14-settings',
 'Site Settings is updated by admins. Its Save action changes settings used throughout the website.',[
 'In Identity, review Site Name, Legal Name, localized Tagline, light and dark logos, Default OG Image, and Social Profiles.',
 'In Contact & Offices, maintain office records, General Email, Notification Emails for enquiries, and Careers Emails for applications.',
 'In Integrations, maintain the Calendly booking URL and approved analytics identifiers. Review Consent Mode with the site owner.',
 'In SEO & Crawlers, review the title suffix, default description, crawler policy, and Search Console verification. Save and check affected public pages.'
 ],'CMS fields do not configure database credentials, email delivery API keys, or storage credentials. A developer maintains those environment settings.'),
('Review applications and contact messages','16-contact',
 'Submissions arrive from the public website. Open existing records in the collection list to review and update them.',[
 'Open Contact messages. Use Filters and Columns to find enquiries by status or date, then open a record and read its details.',
 'Reply using the team inbox or agreed communication channel. Change Status to replied, closed, or spam as appropriate, and Save. Changing status does not send a reply.',
 'Open Applications to review the candidate, Job, CV, and Cover Note. Set status to new, reviewing, interviewing, rejected, or hired, then Save.',
 'Expand Attribution to inspect source page, locale, campaign values, and referrer. Open CVs while logged in; if an expiring download stops working, reopen it from the CMS.'
 ],'Only admins can delete submissions and CV files. Handle candidate and enquiry data according to the organization retention policy.'),
('Manage CMS user access','17-users',
 'Admins create accounts and assign roles. Editors can update their own account details, but cannot promote themselves to admin.',[
 'Open Users and select Create New. Enter the user Email, Name, New Password, and Confirm Password.',
 'Choose Editor for routine content work or Admin for a person who must also manage users and Site Settings.',
 'Save the account and provide access through the team secure process. Each person should use an individual account.',
 'Use the top right account control to manage your own profile and sign out. Use Forgot password on the login page when recovery email is configured.'
 ],'Sessions last 8 hours. Five failed login attempts lock an account for 10 minutes; stop retrying and use the recovery or administrator route.'),
]

def shade(cell, fill):
 p=cell._tc.get_or_add_tcPr(); e=OxmlElement('w:shd');e.set(qn('w:fill'),fill);p.append(e)
def note(doc,text):
 p=doc.add_paragraph();p.paragraph_format.space_before=Pt(10)
 r=p.add_run('Remember  ');r.bold=True;p.add_run(text)
 pr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'EEF7F4');pr.append(sh)
def shot(doc,name,caption,width=6.85):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(4)
 r=p.add_run();r.add_picture(str(ROOT/'screenshots'/f'{name}.png'),width=Inches(width))
 p.paragraph_format.keep_with_next=True
 pic=r._r.xpath('.//wp:docPr')[0];pic.set('descr',caption)
 p=doc.add_paragraph(caption,'Caption');p.paragraph_format.space_after=Pt(10)
def table(doc,headers,rows):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for i,h in enumerate(headers):
  t.rows[0].cells[i].text=h;shade(t.rows[0].cells[i],'E8EEEC')
  for r in t.rows[0].cells[i].paragraphs[0].runs:r.bold=True
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):cells[i].text=v
 for row in t.rows:
  for cell in row.cells:
   for p in cell.paragraphs:
    p.paragraph_format.space_after=Pt(5);p.paragraph_format.space_before=Pt(5)
    for r in p.runs:r.font.size=Pt(10)
 return t

doc=Document();s=doc.sections[0]
s.page_width=Inches(8.5);s.page_height=Inches(11.7)
s.top_margin=Inches(.65);s.bottom_margin=Inches(.6);s.left_margin=Inches(.82);s.right_margin=Inches(.82)
for st in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption','List Number','List Bullet']:
 style=doc.styles[st];style.font.name='Calibri';style.font.color.rgb=RGBColor(0,0,0)
 style.paragraph_format.space_after=Pt(7)
 for border in list(style.element.xpath('.//w:pBdr')): border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(11)
doc.styles['Normal'].paragraph_format.line_spacing=1.08
doc.styles['Title'].font.size=Pt(35);doc.styles['Title'].font.bold=True
doc.styles['Subtitle'].font.size=Pt(15)
doc.styles['Subtitle'].font.italic=False
doc.styles['Heading 1'].font.size=Pt(23);doc.styles['Heading 1'].font.bold=True
doc.styles['Heading 2'].font.size=Pt(14)
doc.styles['Caption'].font.size=Pt(9);doc.styles['Caption'].font.italic=False
doc.core_properties.title='FEKRA CMS User Guide';doc.core_properties.author='FEKRA';doc.core_properties.subject='Content editing and administration'
foot=s.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('FEKRA CMS User Guide   |   ');r.font.size=Pt(9)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)

doc.add_paragraph('FEKRA', 'Subtitle')
doc.add_paragraph('CMS User Guide','Title')
doc.add_paragraph('Content editing and website administration','Subtitle')
doc.add_paragraph('October 2026  |  English edition')
doc.add_paragraph('Use this guide to update the FEKRA website, prepare and publish content, manage translations, and review incoming enquiries and applications. Start by signing in, then follow the workflow for the content you need to change.')
im=Image.open(ROOT/'screenshots/01-login.png');w,h=im.size
im.crop((int(w*.30),int(h*.27),int(w*.70),int(h*.70))).save(ROOT/'screenshots/login-detail.png')
shot(doc,'login-detail','Sign in using your individual CMS account.',3.5)
doc.add_heading('Start with a small update',2)
for t in ['Open https://www.fekra-egy.com/admin and sign in with the account supplied by your administrator.','Choose a collection, open a record, and confirm the editing Locale.','Make the change, save a draft, and review Preview.','Publish approved content and check the public page. Sitewide globals use Save directly.']:
 doc.add_paragraph(t,'List Number')

doc.add_page_break();doc.add_heading('Contents',1)
doc.add_paragraph('Follow a complete workflow or use this page to find a specific task. Page numbers refer to this edition.')
for i,(title,*_) in enumerate(SECTIONS):
 p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
 p.add_run(f'{i+3:02d}   ').bold=True;p.add_run(title)
doc.add_paragraph('21   Publishing checklist and troubleshooting')

md=['# FEKRA CMS User Guide','', 'English edition | October 2026','', 'CMS: https://www.fekra-egy.com/admin','']
for idx,(title,image,intro,steps,reminder) in enumerate(SECTIONS,1):
 doc.add_page_break();doc.add_paragraph(f'WORKFLOW {idx:02d}','Subtitle');doc.add_heading(title,1)
 doc.add_paragraph(intro)
 shot(doc,image,title)
 for i,text in enumerate(steps,1):
  p=doc.add_paragraph();p.paragraph_format.space_after=Pt(8)
  p.add_run(f'{i}.  ').bold=True;p.add_run(text)
 note(doc,reminder)
 md += [f'## {title}','',intro,'',f'![{title}](screenshots/{image}.png)','']+[f'{i}. {v}' for i,v in enumerate(steps,1)]+['',f'**Remember:** {reminder}','']

doc.add_page_break();doc.add_heading('Publishing checklist and troubleshooting',1)
doc.add_heading('Before publishing',2)
checks=['The correct document and locale are selected.','Required fields are complete; copy and translations are approved.','Links lead to the intended published pages and have visible labels.','Images have useful alt text or are correctly marked decorative.','Mobile and desktop previews are readable.','SEO title, description, and social image are set.','Available Locales contains only approved translations.','After publishing, the public page, listing, navigation, and reused sections are checked.']
for t in checks:doc.add_paragraph(t,'List Bullet')
doc.add_heading('When something needs attention',2)
rows=[('Publish is disabled','Wait for the current save to finish. Correct required field errors and check the saved status.'),('Preview cannot open','Save the document first. If it still fails, ask the developer to check preview configuration.'),('A language is missing','Check its localized content and Available Locales. Publish approved changes and check the language switcher.'),('A change is not visible','Confirm the correct record and locale were published. Check whether the section is shared from Home.'),('Translation is disabled','Save the English article changes first, then retry Translate empty languages.'),('A link is invisible','Add a visible Label and a valid destination, then save the parent record.'),('Email or upload fails','Keep the error and time of failure. Ask the administrator or developer to check delivery or storage settings.'),('The CMS shows a server error','Reload once. If it persists, send the developer the page URL and displayed error ID; avoid repeated saves.')]
table(doc,['What you see','What to do'],rows)
note(doc,'For a content correction, review a saved version first. For access problems or repeated server errors, contact your CMS administrator.')
for border in list(doc.element.xpath('.//w:pBdr')): border.getparent().remove(border)
doc.save(ROOT/'FEKRA-CMS-User-Guide.docx')
(ROOT/'CMS-USER-GUIDE.md').write_text('\n'.join(md)+'\n## Before publishing\n\n'+'\n'.join('- '+x for x in checks)+'\n\n## Troubleshooting\n\n'+'\n\n'.join('**'+a+'**  '+b for a,b in rows)+'\n')
(ROOT/'guide-index.json').write_text(json.dumps([{'page':i+3,'title':v[0]} for i,v in enumerate(SECTIONS)],indent=2))
print('Created guide with',len(SECTIONS)+3,'planned pages')
